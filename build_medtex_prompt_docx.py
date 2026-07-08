from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "medtex_site_prompt_2026-07-08.md"
TARGET = ROOT / "Промт_на_сайт_Medtex_2026-07-08.docx"


def set_run_font(run, name="Arial", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def extract_prompt(markdown_text: str) -> str:
    match = re.search(r"```text\n(.*?)\n```", markdown_text, re.S)
    if not match:
        raise ValueError("Prompt code block not found in markdown source.")
    return match.group(1).strip()


def add_styled_paragraph(doc: Document, text: str) -> None:
    if not text.strip():
        doc.add_paragraph("")
        return

    stripped = text.strip()

    if re.match(r"^\d+\.\s", stripped):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(re.sub(r"^\d+\.\s*", "", stripped))
        set_run_font(run)
        return

    if stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(stripped[2:].strip())
        set_run_font(run)
        return

    if stripped.endswith(":"):
        p = doc.add_paragraph()
        run = p.add_run(stripped)
        set_run_font(run, bold=True)
        return

    p = doc.add_paragraph()
    run = p.add_run(stripped)
    set_run_font(run)


def build_doc(source: Path, target: Path) -> None:
    markdown_text = source.read_text(encoding="utf-8")
    prompt_text = extract_prompt(markdown_text)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(48)
    section.bottom_margin = Pt(48)
    section.left_margin = Pt(56)
    section.right_margin = Pt(56)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Промт на сайт Medtex")
    set_run_font(title_run, size=16, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Версия от 8 июля 2026")
    set_run_font(subtitle_run, size=10)

    doc.add_paragraph("")

    for line in prompt_text.splitlines():
        add_styled_paragraph(doc, line)

    doc.save(target)


if __name__ == "__main__":
    build_doc(SOURCE, TARGET)
    print(TARGET)
